"""Actual parser behavior on pinned, read-only card views; no XLSX save or cache."""
from __future__ import annotations

import ast
import copy
import hashlib
import importlib.util
import json
import os
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

if not __debug__:
    raise RuntimeError("Verification requires assertions; do not use -O")

HERE = Path(__file__).resolve().parent
PARSER = Path(os.environ["OBR_MAIN_REMAINING_PARSER"])
BASELINE = Path(os.environ["OBR_MAIN_REMAINING_BASELINE"])
PUBLIC = Path(os.environ["OBR_MAIN_REMAINING_PUBLIC"])
SOURCES = {
    "2014": ("DND5E人物卡_悲灵_弗人_枭熊适配版.xlsx", "94444fda4206d579125418b0007b89c414b86c27d6b98655f81f3f1ce65e8fe6"),
    "2024": ("DND5R人物卡_悲灵_弗人_枭熊适配版.xlsx", "264fc65569e3e80932544af548830ca70f10ee308493c8b0234608fcd01c9b04"),
}
ENGLISH_HEADERS = {
    "L39": "Armor", "P39": "Att.", "R39": "Name", "AL39": "Shield", "AS39": "Worn",
    "L41": "Wond.", "P41": "Att.", "R41": "Rarity", "U41": "Slot", "W41": "Features",
    "L52": "Consum.", "R52": "Rarity", "U52": "Description", "AP52": "Quantity",
    "BJ9": "Class Features", "AX10": "Name", "BC10": "Description",
    "BT2": "Name", "BZ2": "Description", "BT3": "Creature Type", "BT4": "Size", "BT5": "Speed",
}


def module(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    value = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(value)
    return value


class SheetView:
    def __init__(self, source, values, removed=()):
        self.source, self.values = source, values
        self.merged_cells = {str(r) for r in source.merged_cells} - set(removed)

    def __getitem__(self, ref):
        return SimpleNamespace(value=self.values[ref]) if ref in self.values else self.source[ref]

    def cell(self, row, column):
        return self[f"{get_column_letter(column)}{row}"]

    def __getattr__(self, name):
        return getattr(self.source, name)


class BookView:
    def __init__(self, source, values=None, english=False, removed=()):
        self.sheetnames = ["Main" if english and n == "主要" else n for n in source.sheetnames]
        self.sheets = dict(zip(self.sheetnames, [SheetView(source[n], values or {}, removed)
                           if n == "主要" else source[n] for n in source.sheetnames]))

    def __getitem__(self, name):
        return self.sheets[name]


def english_headers(version):
    data = {**ENGLISH_HEADERS, "BT1": "Racial Traits" if version == "2014" else "Species Traits"}
    start = 33 if version == "2014" else 40
    data.update({f"BT{start-2}": "Special Abilities", f"BT{start-1}": "Name", f"BZ{start-1}": "Description"})
    if version == "2024":
        data.update({"BT31": "Fighting Style Feats", "BT32": "Name", "BZ32": "Description"})
    return data


FULL={'2014':['Armor Proficiencies','Weapon Proficiencies','Tool Proficiencies'],
      '2024':['Saving Throw Proficiencies','Skill Proficiencies','Weapon Proficiencies','Tool Proficiencies','Armor Training','Starting Equipment']}
SHORT={'2014':['Armor Prof.','Weapon Prof.','Tool Prof.'],
       '2024':['Save Prof.','Skill Prof.','Weapon Prof.','Tool Prof.','Armor Training','Starting Eq.']}

class RemainingMainTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.before,cls.after=BASELINE.read_bytes(),PARSER.read_bytes()
        assert hashlib.sha256(cls.before).hexdigest()=='2512a21b5d760447887ce323eac3aababcda439e714c71062706186a4e9dce06'
        cls.old,cls.new=module(BASELINE,'remaining_before'),module(PARSER,'remaining_after');cls.books={}
        for v,(name,digest) in SOURCES.items():
            assert hashlib.sha256((PUBLIC/name).read_bytes()).hexdigest()==digest
            cls.books[v]=load_workbook(PUBLIC/name,data_only=True)

    @classmethod
    def tearDownClass(cls):
        for v,book in cls.books.items():
            book.close();name,digest=SOURCES[v];assert hashlib.sha256((PUBLIC/name).read_bytes()).hexdigest()==digest
        assert BASELINE.read_bytes()==cls.before and PARSER.read_bytes()==cls.after

    def view(self,v,mode='short',values=None,removed=()):
        h=english_headers(v);h['BT31' if v=='2014' else 'BT38']='Special Ability'
        h.update({'AX'+str(i+11):s for i,s in enumerate((FULL if mode=='full' else SHORT)[v])})
        return BookView(self.books[v],{**h,**(values or {})},True,removed)

    def parse(self,function,view,old=False):
        m=self.old if old else self.new;return getattr(m,function)(view,m._select_layout(view))

    def test_01_only_two_functions_change(self):
        a,b=ast.parse(self.before),ast.parse(self.after)
        self.assertEqual(len(a.body),len(b.body))
        changed={getattr(x,'name',None) for x,y in zip(a.body,b.body) if ast.dump(x)!=ast.dump(y)}
        self.assertEqual(changed,{'_calibrate_main_fields','parse_special_resources'})

    def test_02_original_full_imports_unchanged(self):
        for v,book in self.books.items():
            results=[]
            for m in [self.old,self.new]:
                with patch.object(m,'load_workbook',return_value=book):r=m.parse_character(PUBLIC/SOURCES[v][0])
                r['meta'].pop('parsed_at');results.append(r)
            self.assertEqual(*results)

    def test_03_short_full_mixed_case_headers_keep_exact_layout(self):
        for v in self.books:
            expected=self.new._select_layout(self.books[v])['main']
            for mode in ['short','full']:
                view=self.view(v,mode)
                self.assertEqual(self.new._select_layout(view)['main'],expected)
                upper={k:'  '+str(value).upper()+'  ' for k,value in view['Main'].values.items()}
                self.assertEqual(self.new._select_layout(self.view(v,mode,upper))['main'],expected)

    def test_04_feature_rows_include_boundaries_exclude_metadata(self):
        for v in self.books:
            first=14 if v=='2014' else 17;special=33 if v=='2014' else 40
            values={f'AX{r}':f'Feature {r}' for r in range(first,78)}
            values.update({f'BT{r}':f'Special {r}' for r in range(special,special+5)})
            values.update({'AX78':'Outside class',f'BT{special+5}':'Outside special','BT8':'Name','BT9':'Speed'})
            for mode in ['short','full']:
                got=self.parse('parse_features',self.view(v,mode,values))
                self.assertEqual([r['name'] for r in got['class_features']],[f'Feature {r}' for r in range(first,78)])
                self.assertEqual([r['name'] for r in got['special_abilities']],[f'Special {r}' for r in range(special,special+5)])
                self.assertIn('Name',[r['name'] for r in got['race_features']]);self.assertIn('Speed',[r['name'] for r in got['race_features']])

    def test_05_nine_default_resources_are_not_real_resources(self):
        for v in self.books:
            for label in ['Special Ability','SPECIAL ABILITY','  special ability  ','特殊能力','名称','',None]:
                values={col+str(row):label for col in ['B','Q','AF'] for row in range(26,29)}
                got=self.parse('parse_special_resources',self.view(v,values=values));self.assertEqual(got,[])
                if label=='Special Ability':self.assertEqual(len(self.parse('parse_special_resources',self.view(v,values=values),old=True)),9)

    def test_06_named_resources_and_values_survive(self):
        for v in self.books:
            layout=self.new._select_layout(self.view(v));cfg=layout['main'];values={};expected=[]
            names=['Name','Speed','Special Ability: Rage','Special Abilities','Ki','Focus','Second Wind','Charges','Custom resource']
            for i,(row,col) in enumerate((row,col) for row in range(26,29) for col in cfg['special_resource_cols']):
                values.update({col['name']+str(row):names[i],col['desc']+str(row):'Resource detail',col['cur']+str(row):i,col['max']+str(row):i+3,col['recharge']+str(row):'Short Rest'})
                expected.append({'name':names[i],'description':'Resource detail','current':i,'max':i+3,'recharge':'Short Rest'})
            view=self.view(v,values=values);self.assertEqual(self.parse('parse_special_resources',view),expected)
            self.assertEqual(self.parse('parse_special_resources',view,old=True),expected)

    def test_07_unknown_headers_or_missing_geometry_keep_previous_fallback(self):
        for v in self.books:
            start=33 if v=='2014' else 40
            for values,removed,key in [({'AX11':'Unrecognized header'},(),'feat_class_rows'),({},('BC77:BR77',),'feat_class_rows'),({f'BT{start-2}':'Unknown'},(),'feat_special_ability_rows'),({},(f'BZ{start+4}:CO{start+4}',),'feat_special_ability_rows')]:
                view=self.view(v,values=values,removed=removed)
                self.assertEqual(self.new._select_layout(view)['main'][key],self.old._select_layout(view)['main'][key])

    def test_08_no_global_or_legacy_layout_mutation(self):
        frozen=copy.deepcopy(self.new.LAYOUT_V1_0_12)
        for v in ['2014','2024','2014']:
            view=self.view(v);self.new._select_layout(view)['main']['feat_class_rows']=(99,100)
            self.assertEqual(self.new._select_layout(view,'v1.0.0'),self.old._select_layout(view,'v1.0.0'))
        self.assertEqual(self.new.LAYOUT_V1_0_12,frozen)

if __name__=='__main__':unittest.main(verbosity=2)
