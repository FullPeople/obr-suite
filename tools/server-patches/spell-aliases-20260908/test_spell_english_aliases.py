"""Real parser + read-only source-card views; never writes or recalculates XLSX."""
from __future__ import annotations

import copy
import hashlib
import importlib.util
import os
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PARSER = Path(os.environ.get("OBR_SPELL_ALIAS_PARSER", Path(__file__).resolve().parent / "parser.py"))
BASELINE = Path(os.environ.get("OBR_SPELL_ALIAS_BASELINE", ROOT / "_audit/2026-09-08/xlsx-server-spell-alias/parser-before-spell-alias.py"))
PUBLIC = Path(os.environ.get("OBR_SPELL_ALIAS_PUBLIC", ROOT / "obr-suite/public"))
spec = importlib.util.spec_from_file_location("spell_alias_parser", PARSER)
parser = importlib.util.module_from_spec(spec)
spec.loader.exec_module(parser)
BEFORE_SHA = "ff02e712b37dd855e1af189969365c6049895d95c93733cb8126fe5610400766"
SOURCES = {
    "2014": ("DND5E人物卡_悲灵_弗人_枭熊适配版.xlsx", "94444fda4206d579125418b0007b89c414b86c27d6b98655f81f3f1ce65e8fe6"),
    "2024": ("DND5R人物卡_悲灵_弗人_枭熊适配版.xlsx", "264fc65569e3e80932544af548830ca70f10ee308493c8b0234608fcd01c9b04"),
}


class SheetView:
    def __init__(self, source, changes):
        self.source, self.changes = source, changes

    def __getitem__(self, ref):
        return SimpleNamespace(value=self.changes[ref]) if ref in self.changes else self.source[ref]

    def cell(self, row, column):
        return self[f"{get_column_letter(column)}{row}"]

    def __getattr__(self, key):
        return getattr(self.source, key)


class BookView:
    def __init__(self, source, changes=None, english=False):
        names = {"主要": "Main", "法术大全": "Spell Compendium"} if english else {}
        self.sheets = {names.get(n, n): SheetView(source[n], (changes or {}).get(n, {})) for n in source.sheetnames}
        self.sheetnames = list(self.sheets)

    def __getitem__(self, name):
        return self.sheets[name]


def stable(data):
    data = copy.deepcopy(data)
    data["meta"].pop("parsed_at")
    return data


class SpellEnglishAliasTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        assert hashlib.sha256(BASELINE.read_bytes()).hexdigest() == BEFORE_SHA
        spec = importlib.util.spec_from_file_location("spell_alias_before", BASELINE)
        cls.old = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.old)
        cls.books, cls.paths = {}, {}
        for version, (filename, digest) in SOURCES.items():
            path = PUBLIC / filename
            assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
            cls.paths[version] = path
            cls.books[version] = load_workbook(path, data_only=True)

    @classmethod
    def tearDownClass(cls):
        for version, book in cls.books.items():
            book.close()
            assert hashlib.sha256(cls.paths[version].read_bytes()).hexdigest() == SOURCES[version][1]
        assert hashlib.sha256(BASELINE.read_bytes()).hexdigest() == BEFORE_SHA

    def full(self, book, version="2014", old=False):
        mod = self.old if old else parser
        with patch.object(mod, "load_workbook", return_value=book):
            return stable(mod.parse_character(self.paths[version]))

    def casting(self, book, old=False):
        mod = self.old if old else parser
        return mod.parse_spellcasting(book, mod._select_layout(book))

    def entries(self, cast):
        return cast["always_known"] + cast["prepared"] + cast["cantrips_known"]

    def selected(self, book, name, old=False):
        return next(e for e in self.entries(self.casting(book, old)) if e["name"] == name.strip())

    def test_01_original_full_character_results_unchanged(self):
        for version, book in self.books.items():
            with self.subTest(version=version):
                self.assertEqual(self.full(book, version), self.full(book, version, old=True))

    def test_02_actual_q77_english_body_matches_same_row_chinese_full_parse(self):
        for version, book in self.books.items():
            with self.subTest(version=version):
                zh = BookView(book, {"主要": {"P77": 0, "Q77": "酸液飞溅"}})
                en = BookView(book, {"主要": {"P77": 0, "Q77": "Acid Splash"}})
                self.assertNotIn("description", self.selected(en, "Acid Splash", old=True))
                before, after = self.full(zh, version, old=True), self.full(en, version)
                row = next(e for e in before["spellcasting"]["cantrips_known"] if e["name"] == "酸液飞溅")
                self.assertEqual(row["description"], book["法术大全"]["M3"].value.strip())
                row["name"] = "Acid Splash"
                self.assertEqual(after, before)

    def test_03_every_existing_spell_group_enriches_without_rewriting_names(self):
        for version, book in self.books.items():
            for lv, name in (("P66", "Q66"), ("X66", "Y66"), ("AF66", "AG66"), ("AN66", "AO66"), ("P77", "Q77"), ("X77", "Y77")):
                with self.subTest(version=version, name=name):
                    zh = BookView(book, {"主要": {lv: 0, name: "酸液飞溅"}})
                    en = BookView(book, {"主要": {lv: 0, name: "Acid Splash"}})
                    expected = self.casting(zh, old=True)
                    next(e for e in self.entries(expected) if e["name"] == "酸液飞溅")["name"] = "Acid Splash"
                    self.assertEqual(self.casting(en), expected)

    def test_04_case_whitespace_and_resolved_english_sheet_roles(self):
        for version, book in self.books.items():
            for name in ("Acid Splash", " acid splash ", "ACID SPLASH"):
                with self.subTest(version=version, name=name):
                    view = BookView(book, {"主要": {"P77": 0, "Q77": name}}, english=True)
                    entry = self.selected(view, name)
                    self.assertEqual(entry["description"], book["法术大全"]["M3"].value.strip())
                    self.assertEqual(entry["name"], name.strip())

    def test_05_real_2024_ambiguous_english_name_is_not_guessed(self):
        book = self.books["2024"]
        self.assertEqual(book["法术大全"]["N452"].value, book["法术大全"]["N646"].value)
        self.assertNotEqual(book["法术大全"]["A452"].value, book["法术大全"]["A646"].value)
        view = BookView(book, {"主要": {"P77": 0, "Q77": "Sanctum of the Flock"}})
        self.assertEqual(self.full(view, "2024"), self.full(view, "2024", old=True))
        self.assertNotIn("description", self.selected(view, "Sanctum of the Flock"))

    def test_06_duplicate_english_even_case_variant_or_same_target_is_rejected(self):
        for version, book in self.books.items():
            for other in ("Acid Splash", " acid splash ", "ACID SPLASH"):
                for target in (book["法术大全"]["A4"].value, "酸液飞溅"):
                    with self.subTest(version=version, other=other, target=target):
                        view = BookView(book, {"主要": {"P77": 0, "Q77": "Acid Splash"},
                                              "法术大全": {"A4": target, "N4": other}})
                        self.assertEqual(self.casting(view), self.casting(view, old=True))

    def test_07_duplicate_canonical_rows_cannot_redirect_an_earlier_alias(self):
        for version, book in self.books.items():
            for name in ("Acid Splash", "Blade Ward", "酸液飞溅"):
                with self.subTest(version=version, name=name):
                    view = BookView(book, {"主要": {"P77": 0, "Q77": name},
                                          "法术大全": {"A4": "酸液飞溅"}})
                    self.assertEqual(self.casting(view), self.casting(view, old=True))

    def test_08_canonical_keys_win_and_folded_collisions_do_not_guess(self):
        for version, book in self.books.items():
            for name in ("Acid Splash", "acid splash", "剑刃防护"):
                with self.subTest(version=version, name=name):
                    view = BookView(book, {"主要": {"P77": 0, "Q77": name},
                                          "法术大全": {"A4": "Acid Splash"}})
                    self.assertEqual(self.casting(view), self.casting(view, old=True))
            view = BookView(book, {"主要": {"P77": 0, "Q77": "剑刃防护"},
                                  "法术大全": {"N3": "剑刃防护"}})
            self.assertEqual(self.casting(view), self.casting(view, old=True))

    def test_09_empty_alias_or_missing_canonical_row_does_not_invent_data(self):
        for version, book in self.books.items():
            for changes in ({"N3": None}, {"N3": " "}, {"A3": None}):
                with self.subTest(version=version, changes=changes):
                    view = BookView(book, {"主要": {"P77": 0, "Q77": "Acid Splash"}, "法术大全": changes})
                    self.assertEqual(self.casting(view), self.casting(view, old=True))

    def test_10_alternating_versions_imports_use_only_the_current_body(self):
        expected = {v: b["法术大全"]["M3"].value.strip() for v, b in self.books.items()}
        self.assertNotEqual(expected["2014"], expected["2024"])
        for version in ("2014", "2024", "2014"):
            view = BookView(self.books[version], {"主要": {"P77": 0, "Q77": "Acid Splash"}})
            self.assertEqual(self.selected(view, "Acid Splash")["description"], expected[version])

    def test_11_legacy_layouts_preserve_existing_results_and_accept_local_unique_alias(self):
        for version in ("v1.0.0", "v1.0.11", "v1.0.12"):
            book = Workbook(); book.active.title = "主要"
            book.create_sheet("背景" if version == "v1.0.0" else "起源")
            db = book.create_sheet("法术大全")
            book["主要"]["A1"] = f"DND card<Test {version}>"
            db["A3"], db["M3"], db["N3"] = "旧版法术", "Legacy-only body", "Legacy Spell"
            layout = parser._select_layout(book)
            row = layout["main"]["spell_cantrip_rows"][0]
            lvcol, namecol = layout["main"]["spell_cantrip_cols"][0]
            book["主要"][f"{lvcol}{row}"] = 0
            book["主要"][f"{namecol}{row}"] = "旧版法术"
            with self.subTest(version=version):
                self.assertEqual(self.full(book), self.full(book, old=True))
                book["主要"][f"{namecol}{row}"] = "Legacy Spell"
                self.assertEqual(self.selected(book, "Legacy Spell")["description"], "Legacy-only body")
            book.close()

    def test_12_database_original_keys_values_and_length_are_unchanged(self):
        for version, book in self.books.items():
            self.assertEqual(parser.load_spell_db(book, parser._select_layout(book)),
                             self.old.load_spell_db(book, self.old._select_layout(book)))

    def test_13_unknown_names_and_existing_header_filter_remain_unchanged(self):
        for version, book in self.books.items():
            for name in ("Unknown custom spell", "Spell Name", "法术名称"):
                view = BookView(book, {"主要": {"P77": 0, "Q77": name}})
                self.assertEqual(self.full(view, version), self.full(view, version, old=True))

    def test_14_cached_level_and_missing_body_are_not_recalculated_or_replaced(self):
        for version, book in self.books.items():
            for level in (None, 7):
                zh = BookView(book, {"主要": {"P77": level, "Q77": "酸液飞溅"}, "法术大全": {"M3": None}})
                en = BookView(book, {"主要": {"P77": level, "Q77": "Acid Splash"}, "法术大全": {"M3": None}})
                expected = self.selected(zh, "酸液飞溅", old=True)
                expected["name"] = "Acid Splash"
                self.assertEqual(self.selected(en, "Acid Splash"), expected)


if __name__ == "__main__":
    unittest.main(verbosity=2)
