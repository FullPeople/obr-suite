"""Actual parser behavior on pinned, read-only card views; no XLSX save or cache."""
from __future__ import annotations

import ast
import copy
import hashlib
import importlib.util
import json
import os
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

if not __debug__:
    raise RuntimeError("Verification requires assertions; do not use -O")

HERE = Path(__file__).resolve().parent
PARSER = Path(os.environ["OBR_MAIN_FIELDS_PARSER"])
BASELINE = Path(os.environ["OBR_MAIN_FIELDS_BASELINE"])
PUBLIC = Path(os.environ["OBR_MAIN_FIELDS_PUBLIC"])
SOURCES = {
    "2014": ("DND5E人物卡_悲灵_弗人_枭熊适配版.xlsx", "94444fda4206d579125418b0007b89c414b86c27d6b98655f81f3f1ce65e8fe6"),
    "2024": ("DND5R人物卡_悲灵_弗人_枭熊适配版.xlsx", "264fc65569e3e80932544af548830ca70f10ee308493c8b0234608fcd01c9b04"),
}
ENGLISH_HEADERS = {
    "L39": "Armor", "P39": "Att.", "R39": "Name", "AL39": "Shield", "AS39": "Worn",
    "L41": "Wond.", "P41": "Att.", "R41": "Rarity", "U41": "Slot", "W41": "Features",
    "L52": "Consum.", "R52": "Rarity", "U52": "Description", "AP52": "Quantity",
    "BJ9": "Class Features", "AX10": "Name", "BC10": "Description",
    "BT2": "Name", "BZ2": "Description", "BT3": "Creature Type", "BT4": "Size", "BT5": "Speed",
}


def module(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    value = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(value)
    return value


class SheetView:
    def __init__(self, source, values, removed=()):
        self.source, self.values = source, values
        self.merged_cells = {str(r) for r in source.merged_cells} - set(removed)

    def __getitem__(self, ref):
        return SimpleNamespace(value=self.values[ref]) if ref in self.values else self.source[ref]

    def cell(self, row, column):
        return self[f"{get_column_letter(column)}{row}"]

    def __getattr__(self, name):
        return getattr(self.source, name)


class BookView:
    def __init__(self, source, values=None, english=False, removed=()):
        self.sheetnames = ["Main" if english and n == "主要" else n for n in source.sheetnames]
        self.sheets = dict(zip(self.sheetnames, [SheetView(source[n], values or {}, removed)
                           if n == "主要" else source[n] for n in source.sheetnames]))

    def __getitem__(self, name):
        return self.sheets[name]


def english_headers(version):
    data = {**ENGLISH_HEADERS, "BT1": "Racial Traits" if version == "2014" else "Species Traits"}
    start = 33 if version == "2014" else 40
    data.update({f"BT{start-2}": "Special Abilities", f"BT{start-1}": "Name", f"BZ{start-1}": "Description"})
    if version == "2024":
        data.update({"BT31": "Fighting Style Feats", "BT32": "Name", "BZ32": "Description"})
    return data


class MainFieldsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.before_bytes, cls.after_bytes = BASELINE.read_bytes(), PARSER.read_bytes()
        assert hashlib.sha256(cls.before_bytes).hexdigest() == "a07204634e7f6b5a306ed444c6fbc5ef699c53a2b95331f6f2de3b729de7febe"
        cls.old, cls.new = module(BASELINE, "main_fields_before"), module(PARSER, "main_fields_after")
        cls.books = {}
        for version, (name, digest) in SOURCES.items():
            assert hashlib.sha256((PUBLIC/name).read_bytes()).hexdigest() == digest
            cls.books[version] = load_workbook(PUBLIC/name, data_only=True)

    @classmethod
    def tearDownClass(cls):
        for version, book in cls.books.items():
            book.close()
            name, digest = SOURCES[version]
            assert hashlib.sha256((PUBLIC/name).read_bytes()).hexdigest() == digest
        assert BASELINE.read_bytes() == cls.before_bytes and PARSER.read_bytes() == cls.after_bytes

    def view(self, version, values=None, english=False, removed=()):
        return BookView(self.books[version], {**(english_headers(version) if english else {}), **(values or {})}, english, removed)

    def parse(self, function, view, old=False):
        mod = self.old if old else self.new
        return getattr(mod, function)(view, mod._select_layout(view))

    def full(self, view, version, old=False):
        mod = self.old if old else self.new
        with patch.object(mod, "load_workbook", return_value=view):
            data = mod.parse_character(PUBLIC/SOURCES[version][0])
        data["meta"].pop("parsed_at")
        return data

    def test_01_all_sections_resolve_actual_two_cards_and_short_english_headers(self):
        for version in self.books:
            for english in (False, True):
                with self.subTest(version=version, english=english):
                    fields = self.new._select_layout(self.view(version, english=english))["main"]
                    self.assertEqual(fields["armor_name"], "L40")
                    self.assertIsNone(fields["wondrous_qty_col"])
                    self.assertEqual(fields["consumable_qty_col"], "AP")
                    self.assertEqual([fields["currency_"+c] for c in ("gp","pp","ep","sp","cp")], ["AI59","AP59","AI60","AI61","AI62"])
                    self.assertEqual(fields["feat_class_rows"], (14 if version=="2014" else 17,78))
                    self.assertEqual(fields["feat_race_rows"], (6,16))
                    self.assertEqual(fields["feat_special_ability_rows"], (33,38) if version=="2014" else (40,45))
                    self.assertEqual(fields["feat_fighting_style_rows"], None if version=="2014" else (33,38))

    def test_02_equipped_short_caption_does_not_fall_back_to_as41(self):
        for version in self.books:
            for value, expected in (("Yes",True),("是",True),("No",False),("否",False),(None,False)):
                view=self.view(version,{"AS40":value,"AS41":"Yes","AS39":" worn "},True)
                self.assertEqual(self.parse("parse_combat",view)["shield"]["equipped"],expected)

    def test_03_wondrous_ten_rows_do_not_read_adjacent_skills(self):
        for version in self.books:
            values={f"{col}{r}":value for r in range(42,52) for col,value in (("L",f"Item {r}"),("W",f"Text {r}"),("I",700+r))}
            got=self.parse("parse_wondrous_items",self.view(version,values,True))
            self.assertEqual([(r["name"],r["quantity"]) for r in got],[(f"Item {r}",1) for r in range(42,52)])
            values.update({"L45":None,"L51":"  ","L52":"Consum."})
            self.assertEqual(len(self.parse("parse_wondrous_items",self.view(version,values,True))),8)

    def test_04_consumable_left_quantity_and_last_row(self):
        for version in self.books:
            values={f"{col}{r}":value for r in range(53,58) for col,value in (("L",f"Potion {r}"),("AP",r-50),("AS",99),("I",888))}
            got=self.parse("parse_consumables",self.view(version,values,True))
            self.assertEqual([(r["name"],r["quantity"]) for r in got],[(f"Potion {r}",r-50) for r in range(53,58)])
            values["L55"]=None
            self.assertEqual(len(self.parse("parse_consumables",self.view(version,values))),4)

    def test_05_currency_and_existing_numeric_policy(self):
        for version in self.books:
            for vals in ((13,29,41,53,67),(0,None,"9",-7,1.5),("",True,False,"2 coins","bad")):
                values=dict(zip(("AI59","AP59","AI60","AI61","AI62"),vals))
                values["AN61"]="999GP"
                currency=self.parse("parse_inventory",self.view(version,values))["currency"]
                self.assertEqual(currency["wallet"],dict(zip(("gp","pp","ep","sp","cp"),[self.old._int(v,default=0) for v in vals])))
                self.assertEqual(currency["total_gp_raw"],"999GP")

    def test_06_special_abilities_do_not_read_familiar_or_reclassify_2024_styles(self):
        for version in self.books:
            start=33 if version=="2014" else 40
            values={f"{col}{r}":f"{col} {r}" for r in range(start,start+5) for col in ("BT","BZ")}
            values[f"BT{start+5}"]="Outside special abilities"
            if version=="2024":values.update({f"BT{r}":f"Style {r}" for r in range(33,38)})
            got=self.parse("parse_features",self.view(version,values,True))
            self.assertEqual([r["name"] for r in got["special_abilities"]],[f"BT {r}" for r in range(start,start+5)])
            self.assertEqual([r["name"] for r in got["fighting_style_feats"]],[] if version=="2014" else [f"Style {r}" for r in range(33,38)])
            values[f"BT{start+2}"]=None
            self.assertEqual(len(self.parse("parse_features",self.view(version,values))["special_abilities"]),4)

    def test_07_all_class_rows_including_late_rows_and_no_metadata(self):
        for version in self.books:
            start=14 if version=="2014" else 17
            values={f"{col}{r}":value for r in range(start,78) for col,value in (("AX",f"Feature {r}"),("BC",f"Detail {r}"),("AW",r))}
            values["AX78"]="Outside class table"
            got=self.parse("parse_features",self.view(version,values,True))["class_features"]
            self.assertEqual(got,[{"name":f"Feature {r}","description":f"Detail {r}","level":r} for r in range(start,78)])
            values.update({"AX37":None,"AX55":"  "})
            self.assertEqual(len(self.parse("parse_features",self.view(version,values))["class_features"]),78-start-2)

    def test_08_traits_exclude_bilingual_metadata_and_keep_player_names(self):
        for version in self.books:
            values={f"BT{r}":f"Trait {r}" for r in range(6,16)}
            values.update({"BT16":"Outside traits","BT8":"Name","BT9":"Speed"})
            got=self.parse("parse_features",self.view(version,values,True))["race_features"]
            self.assertEqual([r["name"] for r in got],[values[f"BT{r}"] for r in range(6,16)])

    def test_09_armor_uses_selected_item_and_preserves_other_combat_data(self):
        for version in self.books:
            view=self.view(version,{"L40":"  Plate  ","H40":"wrong skill","R40":"My Armor"})
            old=self.parse("parse_combat",view,True);new=self.parse("parse_combat",view)
            self.assertEqual(new["armor"]["name"],"Plate")
            old["armor"]["name"]="Plate"
            self.assertEqual(new,old)

    def test_10_missing_header_or_geometry_only_retains_that_section(self):
        cases=[("L39","L40:O40","armor_name"),("L41","W51:AT51","wondrous_qty_col"),
               ("AP52","AP57:AQ57","consumable_qty_col"),("AG59","AI62:AM62","currency_gp"),
               ("BJ9","BC77:BR77","feat_class_rows"),("BT1","BZ15:CO15","feat_race_rows")]
        for version in self.books:
            start=33 if version=="2014" else 40
            for header,region,key in cases+[(f"BT{start-2}",f"BZ{start+4}:CO{start+4}","feat_special_ability_rows")]:
                for values,removed in (({header:"Unknown section"},()),({},(region,))):
                    with self.subTest(version=version,key=key,removed=removed):
                        view=self.view(version,values,removed=removed)
                        old=self.old._select_layout(view)["main"];new=self.new._select_layout(view)["main"]
                        self.assertEqual(new[key],old[key])
                        # A different independently proven section still calibrates.
                        self.assertEqual(new["consumable_qty_col" if key=="currency_gp" else "currency_gp"],"AP" if key=="currency_gp" else "AI59")

    def test_11_layouts_are_per_import_and_legacy_override_unchanged(self):
        modern=copy.deepcopy(self.new.LAYOUT_V1_0_12);legacy=copy.deepcopy(self.new.LAYOUT_V1_0_0)
        for version in ("2014","2024","2014","2024"):
            view=self.view(version)
            self.new._select_layout(view)["main"]["armor_name"]="mutated local copy"
            self.assertEqual(self.new._select_layout(view)["main"]["armor_name"],"L40")
            self.assertEqual(self.new._select_layout(view,"v1.0.0"),self.old._select_layout(view,"v1.0.0"))
        self.assertEqual(self.new.LAYOUT_V1_0_12,modern);self.assertEqual(self.new.LAYOUT_V1_0_0,legacy)

    def test_12_full_import_changes_only_explicit_expected_fields(self):
        for version in self.books:
            view=self.view(version,{"AV1":None,"L40":"Plate","AI59":13,"AP59":29,"AI60":41,"AI61":53,"AI62":67,
                                    "L42":"Wand","L53":"Potion","AP53":4,"AX77":"Late feature"})
            old=self.full(view,version,True);new=self.full(view,version)
            self.assertEqual(new["combat"]["armor"]["name"],"Plate")
            old["combat"]["armor"]["name"]="Plate"
            old["inventory"]["currency"]["wallet"]={"gp":13,"pp":29,"ep":41,"sp":53,"cp":67}
            old["inventory"]["wondrous_items"][0]["quantity"]=1
            old["inventory"]["consumables"][0]["quantity"]=4
            old["features"]["class_features"].append({"name":"Late feature","description":None,"level":None})
            if version=="2014":
                old["features"]["fighting_style_feats"]=[]
                old["features"]["special_abilities"]=[]
            self.assertEqual(new,old)

    def test_13_embedded_json_priority_and_original_cached_import(self):
        for version in self.books:
            # Coordinate corrections must not silently change the existing overlay contract.
            original=self.full(self.view(version),version)
            self.assertIsInstance(original,dict)
            payload={"hp":123,"abilities":{"str":18},"name":"Cached display name",
                     "inventory":{"currency":{"wallet":{"gp":777}},"wondrous_items":[{"name":"Cached wand","quantity":9}]}}
            view=self.view(version,{"AV1":json.dumps(payload),"AI59":13,"L42":"Input wand"})
            new=self.full(view,version)
            old=self.full(view,version,True)
            self.assertEqual(new["core_stats"],old["core_stats"])
            self.assertEqual(new["core_stats"]["hp"]["current"],123)
            self.assertEqual(new["abilities"]["str"]["total"],18)
            self.assertEqual(new["identity"]["display_name"],"Cached display name")
            # Inventory is NOT an accepted overlay key in this parser. The
            # corrected source fields must remain authoritative here.
            self.assertEqual(new["inventory"]["currency"]["wallet"]["gp"],13)
            self.assertEqual(new["inventory"]["wondrous_items"][0]["name"],"Input wand")
            self.assertEqual(new["inventory"]["wondrous_items"][0]["quantity"],1)

    def test_14_unrelated_parser_functions_and_constants_are_unchanged(self):
        def nodes(raw):
            return {n.name:ast.dump(n,include_attributes=False) for n in ast.parse(raw).body if isinstance(n,(ast.FunctionDef,ast.ClassDef))}
        before=nodes(self.before_bytes);after=nodes(self.after_bytes)
        changed={k for k in before if before[k]!=after.get(k)}
        self.assertEqual(changed,{"_select_layout","parse_wondrous_items"})
        self.assertEqual(set(after)-set(before),{"_calibrate_main_fields"})
        other=lambda raw:[ast.dump(n,include_attributes=False) for n in ast.parse(raw).body if not isinstance(n,(ast.FunctionDef,ast.ClassDef))]
        self.assertEqual(other(self.before_bytes),other(self.after_bytes))


if __name__ == "__main__":
    unittest.main(verbosity=2)
