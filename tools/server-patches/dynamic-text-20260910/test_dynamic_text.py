"""Native saved files plus explicit-override and compiler-recognition boundaries."""
import ast,copy,hashlib,importlib.util,json,os,time,unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
if not __debug__:raise RuntimeError('Verification requires assertions')
HERE=Path(__file__).resolve().parent
def module(path,name):
    spec=importlib.util.spec_from_file_location(name,path);result=importlib.util.module_from_spec(spec);spec.loader.exec_module(result);return result
PARSER=Path(os.environ['OBR_DYNAMIC_PARSER']);BASELINE=Path(os.environ['OBR_DYNAMIC_BASELINE']);SAVED=Path(os.environ['OBR_DYNAMIC_SAVED_RESULT'])
new=module(PARSER,'dynamic_parser');old=module(BASELINE,'dynamic_baseline');native=json.loads(SAVED.read_bytes())
class FormulaBook:
    def __init__(self,formula,kind='f'):self.formula=formula;self.kind=kind;self.closed=False
    def __getitem__(self,name):return {'AV1':SimpleNamespace(value=self.formula,data_type=self.kind)}
    def close(self):self.closed=True
class DynamicTextTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cards={};cls.formulas={};cls.raw={};cls.reports={}
        for version,row in native['versions'].items():
            path=Path(row['savedPath']);assert hashlib.sha256(path.read_bytes()).hexdigest()==row['savedSha256'];cls.cards[version]=path
            wb=load_workbook(path,read_only=True,data_only=False);cls.formulas[version]=wb['主要']['AV1'].value;wb.close()
            wb=load_workbook(path,read_only=True,data_only=True);cls.raw[version]={r:wb['主要'][r].value for r in row['expectedInputs']};wb.close()

    def test_01_native_saved_inputs_are_complete(self):
        for version,row in native['versions'].items():self.assertEqual(self.raw[version],row['expectedInputs'])

    def test_02_generated_expression_recognition(self):
        for version,path in self.cards.items():
            ej=native['versions'][version]['afterReopen']['export']
            self.assertTrue(new._is_compiled_template_export(path,'主要',ej))
            aliases=self.formulas[version].replace('Export!$A$',"'Export'!A")
            fake=FormulaBook(aliases)
            with patch.object(new,'load_workbook',return_value=fake):self.assertTrue(new._is_compiled_template_export(path,'主要',ej))
            self.assertTrue(fake.closed)

    def test_03_handwritten_or_different_formulas_keep_overrides(self):
        for version,path in self.cards.items():
            ej=native['versions'][version]['afterReopen']['export'];formula=self.formulas[version]
            for value,kind in [(json.dumps(ej),'s'),('=A1','f'),(formula.replace('32767','32000'),'f'),(formula.replace('Export!$A$3','Export!$A$4'),'f'),('=IF(','f')]:
                fake=FormulaBook(value,kind)
                with patch.object(new,'load_workbook',return_value=fake):self.assertFalse(new._is_compiled_template_export(path,'主要',ej))
                self.assertTrue(fake.closed)
            for value in [{},{'schema':'another'},{'schema':'obr-suite-card/v1','meta':[]},{'schema':'obr-suite-card/v1','meta':{'ruleset':'5E2030'}},{'schema':'obr-suite-card/v1','meta':{'ruleset':[]}},{'schema':'obr-suite-card/v1','meta':{'ruleset':{}}}]:
                with patch.object(new,'load_workbook',side_effect=AssertionError('Should not read workbook')):self.assertFalse(new._is_compiled_template_export(path,'主要',value))

    def test_04_explicit_json_overlay_is_unchanged(self):
        base={'identity':{'character_name':'Sheet name','display_name':'Sheet display','race':{'name':'Sheet race'}},'classes':[{'name':'Wizard'}],'abilities':{'str':{'total':10,'save':{'bonus':0}}},'core_stats':{'hp':{'current':1}},'skills':[{'name':'运动','total':0}]}
        for ej in [{'name':'JSON\nName\tTab','race':'Elf','class':'Rogue','level':4,'str':16,'save_str':3,'hp':20},{'identity':{'display_name':'Displayed','character_name':'Ignored'},'abilities':{'str':12},'skills':{'运动':7}},{'identity':{'character_name':'Only JSON character'}},{'name':'  trimmed  ','hp':0}]:
            a=copy.deepcopy(base);b=copy.deepcopy(base);old._overlay_embedded_json(a,ej);new._overlay_embedded_json(b,ej);self.assertEqual(a,b)

    def test_05_source_identity_retains_numeric_overlay(self):
        data={'identity':{'character_name':'Full\nName\tTab','display_name':'Alias\nHere','race':{'name':'Full\nRace'}},'classes':[{'name':'Full\nClass'}],'abilities':{'str':{'total':10,'save':{'bonus':0}}},'core_stats':{'hp':{'current':1}},'skills':[]}
        identity=copy.deepcopy(data['identity']);classes=copy.deepcopy(data['classes'])
        new._overlay_embedded_json(data,{'name':'flattened','race':'flattened','class':'flattened','str':16,'save_str':3,'hp':20,'level':4},preserve_sheet_identity=True)
        self.assertEqual(data['identity'],identity);self.assertEqual(data['classes'],classes);self.assertEqual(data['abilities']['str']['total'],16);self.assertEqual(data['abilities']['str']['save']['bonus'],3);self.assertEqual(data['core_stats']['hp']['current'],20);self.assertEqual(data['total_level'],4)

    def test_06_actual_saved_card_imports(self):
        for version,path in self.cards.items():
            started=time.perf_counter();a=old.parse_character(path);b=new.parse_character(path);elapsed=time.perf_counter()-started
            expected=native['versions'][version]['expectedInputs'];actualname=b['identity']['display_name'] or b['identity']['character_name']
            self.assertEqual(actualname,expected['E3']);self.assertNotEqual(a['identity']['display_name'],actualname)
            self.assertEqual(b['inventory']['wondrous_items'][0]['properties'],expected['W42'])
            self.assertEqual(b['inventory']['consumables'][0]['description'],expected['U53'])
            self.assertEqual(b['features']['special_abilities'][0]['description'],expected['BZ'+('33' if version=='2014' else '40')])
            a['meta']['parsed_at']=b['meta']['parsed_at'];a['identity']['display_name']=b['identity']['display_name'];self.assertEqual(a,b)
            self.reports[version]={'importPairSeconds':elapsed,'savedSha256':hashlib.sha256(path.read_bytes()).hexdigest(),'displayName':actualname,'sourceBodiesPreserved':True}
        Path(os.environ['OBR_DYNAMIC_NATIVE_REPORT']).write_text(json.dumps(self.reports,ensure_ascii=False,indent=2),encoding='utf8')

    def test_07_only_intended_parser_functions_change(self):
        def functions(path):return {n.name:ast.dump(n,include_attributes=False) for n in ast.parse(path.read_text(encoding='utf8')).body if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef))}
        a=functions(BASELINE);b=functions(PARSER);self.assertEqual(set(b)-set(a),{'_is_compiled_template_export'})
        self.assertEqual({k for k in a if a[k]!=b[k]},{'parse_character','_overlay_embedded_json'})

if __name__=='__main__':unittest.main(verbosity=2)
