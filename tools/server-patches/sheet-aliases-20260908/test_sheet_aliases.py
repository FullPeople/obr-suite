"""Read-only source-workbook/legacy fixtures for versioned sheet-role aliases.

Run: python -B -X utf8 -m unittest -v test_sheet_aliases test_auto_resources
Original files are never saved. Renames and seeded input cells are view-only.
"""
from __future__ import annotations

import copy
import hashlib
import json
import os
from pathlib import Path
import runpy
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import parser

ROOT = Path(__file__).resolve().parent.parent
PUBLIC = Path(os.environ.get("OBR_ALIAS_PUBLIC", ROOT / "obr-suite" / "public"))
ORIGINAL_PARSER = Path(os.environ.get(
    "OBR_ALIAS_BASELINE", ROOT / "_audit/2026-09-08/xlsx-server-aliases/original/parser.py"
))
COMMON = {
    "主要": "Main", "数据表": "Data", "职业": "Classes", "装备": "Equipment",
    "魔宠": "Familiars", "背包": "Inventory", "盟友与魔宠": "Allies and Familiars",
    "法术书": "Spellbook", "法术大全": "Spell Compendium",
    "圆形效应范围": "Areas of Effect", "网页导入": "Web Import", "更新": "Changelog",
}
ENGLISH = {
    "2014": {**COMMON, "背景": "Background", "背景数据": "Background Data",
             "种族": "Races", "据点": "Strongholds", "额外法术": "Additional Spells"},
    "2024": {**COMMON, "起源": "Origin", "背景": "Background Data",
             "种族": "Species", "据点": "Bastions", "专长与据点": "Feats and Bastions",
             "自定义调整栏": "Custom Adjustments"},
}
SOURCE_HASHES = {
    "2014": "94444fda4206d579125418b0007b89c414b86c27d6b98655f81f3f1ce65e8fe6",
    "2024": "264fc65569e3e80932544af548830ca70f10ee308493c8b0234608fcd01c9b04",
}


class SheetView:
    def __init__(self, source, values):
        self._source, self._values = source, values

    def __getitem__(self, key):
        return SimpleNamespace(value=self._values[key]) if key in self._values else self._source[key]

    def cell(self, row, column):
        from openpyxl.utils.cell import get_column_letter
        return self[f"{get_column_letter(column)}{row}"]

    def __getattr__(self, key):
        return getattr(self._source, key)


class WorkbookView:
    def __init__(self, source, names=None, omit=(), extra=None, values=None):
        names, extra, values = names or {}, extra or {}, values or {}
        self._sheets = {
            names.get(name, name): SheetView(source[name], values.get(name, {}))
            for name in source.sheetnames if name not in omit
        }
        for name, original in extra.items():
            self._sheets[name] = SheetView(source[original], values.get(original, {}))
        self.sheetnames = list(self._sheets)

    def __getitem__(self, name):
        return self._sheets[name]


def normalized(data):
    data = copy.deepcopy(data)
    data["meta"].pop("parsed_at")  # Sole volatile value; source_file is compared.
    return data


class SheetAliasTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if hashlib.sha256(ORIGINAL_PARSER.read_bytes()).hexdigest() != "b7f73c7fc06e11679b696b7b08e528ae9aaf455518a4f03bcb2eb1c645b6c9c1":
            raise AssertionError("Prechange parser baseline hash does not match")
        cls.original = runpy.run_path(str(ORIGINAL_PARSER), run_name="original_parser")
        cls.books, cls.paths = {}, {}
        for version, pattern in (("2014", "DND5E*.xlsx"), ("2024", "DND5R*.xlsx")):
            path = next(PUBLIC.glob(pattern))
            actual_hash = hashlib.sha256(path.read_bytes()).hexdigest()
            if actual_hash != SOURCE_HASHES[version]:
                raise AssertionError(f"Source workbook changed: {version}: {actual_hash}")
            cls.paths[version] = path
            cls.books[version] = load_workbook(path, data_only=True)

    @classmethod
    def tearDownClass(cls):
        for version, book in cls.books.items():
            book.close()
            assert hashlib.sha256(cls.paths[version].read_bytes()).hexdigest() == SOURCE_HASHES[version]

    def parsed(self, book, version="2014", original=False):
        if original:
            with patch.dict(self.original["parse_character"].__globals__, {"load_workbook": lambda *a, **k: book}):
                return normalized(self.original["parse_character"](self.paths[version]))
        with patch.object(parser, "load_workbook", return_value=book):
            return normalized(parser.parse_character(self.paths[version]))

    def test_original_chinese_full_result_matches_prechange_parser(self):
        for version, book in self.books.items():
            with self.subTest(version=version):
                self.assertEqual(self.parsed(book, version), self.parsed(book, version, original=True))

    def test_every_reviewed_english_sheet_name_preserves_full_result_and_spell_database(self):
        for version, book in self.books.items():
            with self.subTest(version=version):
                english = WorkbookView(book, ENGLISH[version])
                self.assertEqual(self.parsed(english, version), self.parsed(book, version, original=True))
                layout = parser._select_layout(english)
                self.assertEqual(len(parser.load_spell_db(english, layout)), 522 if version == "2014" else 808)
                self.assertEqual(layout["sheets"]["main"], "Main")
                self.assertEqual(layout["sheets"]["inventory"], "Inventory")

    def test_partial_unique_aliases_and_origin_plural_are_allowed(self):
        for version, book in self.books.items():
            for names in ({"主要": "Main"}, {"背包": "Inventory", "法术大全": "Spell Compendium"},
                          {"背景": "Background"} if version == "2014" else {"起源": "Origins"}):
                with self.subTest(version=version, names=names):
                    self.assertEqual(self.parsed(WorkbookView(book, names), version), self.parsed(book, version, original=True))

    def test_actual_spell_text_inventory_and_character_numbers_survive_renaming(self):
        for version, book in self.books.items():
            spell = book["法术大全"]["A3"].value
            av1 = json.loads(book["主要"]["AV1"].value)
            av1["core_stats"]["ac"] = "18"  # Match edited input; retain existing overlay precedence.
            values = {"主要": {"E3": "Alias Test", "E4": "Player", "F13": 18,
                                "R22": 7, "V22": 21, "T24": 3, "D23": 18,
                                "X66": 0, "Y66": spell, "AV1": json.dumps(av1, ensure_ascii=False)},
                      "背包": {"B5": "Test Rope", "G5": "Do not lose this item.", "S5": 10, "V5": 2}}
            before = WorkbookView(book, values=values)
            after = WorkbookView(book, ENGLISH[version], values=values)
            parsed = self.parsed(after, version)
            self.assertEqual(parsed, self.parsed(before, version, original=True))
            self.assertEqual(parsed["abilities"]["str"]["total"], 18)
            self.assertEqual(parsed["core_stats"]["hp"], {"current": 7, "max": 21, "temp": 3})
            self.assertEqual(parsed["core_stats"]["ac"], 18)
            prepared = next(s for s in parsed["spellcasting"]["prepared"] if s["name"] == spell)
            self.assertEqual(prepared["description"], book["法术大全"]["M3"].value.strip())
            self.assertIn("Test Rope", json.dumps(parsed["inventory"], ensure_ascii=False))
            self.assertIn("Do not lose this item.", json.dumps(parsed["inventory"], ensure_ascii=False))

    def test_published_card_missing_critical_sheet_is_rejected(self):
        for version, book in self.books.items():
            person = "背景" if version == "2014" else "起源"
            database = "背景数据" if version == "2014" else "背景"
            for missing in ("主要", person, database, "法术大全", "背包", "装备", "数据表", "圆形效应范围", "网页导入"):
                with self.subTest(version=version, missing=missing), self.assertRaises(ValueError):
                    self.parsed(WorkbookView(book, omit=(missing,)), version)

    def test_same_role_bilingual_duplicates_are_rejected(self):
        for version, book in self.books.items():
            for old, alias in ENGLISH[version].items():
                with self.subTest(version=version, role=old), self.assertRaisesRegex(ValueError, "Ambiguous"):
                    self.parsed(WorkbookView(book, extra={alias: old}), version)
        with self.assertRaisesRegex(ValueError, "Ambiguous"):
            self.parsed(WorkbookView(self.books["2024"], {"起源": "Origin"}, extra={"Origins": "起源"}), "2024")

    def test_cross_version_personal_or_database_names_are_rejected(self):
        with self.assertRaises(ValueError):
            self.parsed(WorkbookView(self.books["2014"], extra={"Origin": "背景"}))
        for rename in ({"起源": "Background"}, {"背景": "背景数据"}):
            with self.subTest(rename=rename), self.assertRaises(ValueError):
                self.parsed(WorkbookView(self.books["2024"], rename), "2024")

    def test_conflicting_or_unsupported_ruleset_markers_are_rejected(self):
        for version, book in self.books.items():
            opposite = "5E2024" if version == "2014" else "5E2014"
            for marker in (opposite, "5E2014 5E2024", "5E2025"):
                with self.subTest(version=version, marker=marker), self.assertRaises(ValueError):
                    self.parsed(WorkbookView(book, values={"主要": {"E2": marker}}), version)
            invalid = {"schema": "obr-suite-card/v1", "meta": {"ruleset": "5E2025"}}
            with self.assertRaises(ValueError):
                self.parsed(WorkbookView(book, values={"主要": {"AV1": json.dumps(invalid)}}), version)

    def test_other_version_reserved_titles_cannot_hide_same_role_duplicates(self):
        for version, wrong in (("2014", {"种族": "Species", "据点": "Bastions"}),
                               ("2024", {"种族": "Races", "据点": "Strongholds"})):
            for name, alias in wrong.items():
                with self.subTest(version=version, alias=alias):
                    with self.assertRaisesRegex(ValueError, "conflicts"):
                        self.parsed(WorkbookView(self.books[version], {name: alias}), version)
                    with self.assertRaisesRegex(ValueError, "Ambiguous"):
                        self.parsed(WorkbookView(self.books[version], extra={alias: name}), version)

    def test_custom_or_legacy_av1_overlay_is_unchanged(self):
        for version, book in self.books.items():
            for av1 in ('{"hp": 12, "abilities": {"str": 16}}', ".st str:16", "bad JSON"):
                values = {"主要": {"AV1": av1}}
                with self.subTest(version=version, av1=av1):
                    self.assertEqual(self.parsed(WorkbookView(book, ENGLISH[version], values=values), version),
                                     self.parsed(WorkbookView(book, values=values), version, original=True))

    def test_verified_av1_can_supply_layout_evidence_without_changing_ruleset_output(self):
        for version, book in self.books.items():
            values = {"主要": {"E2": None}}
            actual = self.parsed(WorkbookView(book, ENGLISH[version], values=values), version)
            self.assertEqual(actual, self.parsed(WorkbookView(book, values=values), version, original=True))
            self.assertIsNone(actual["meta"]["ruleset"])

    def test_markerless_unversioned_main_is_not_a_supported_template(self):
        for version, book in self.books.items():
            with self.subTest(version=version), self.assertRaisesRegex(ValueError, "version/ruleset"):
                self.parsed(WorkbookView(book, ENGLISH[version], values={"主要": {"E2": None, "AV1": None}}), version)

    def test_legacy_coordinate_paths_remain_compatible(self):
        for ver in ("v1.0.0", "v1.0.11", "v1.0.12"):
            book = Workbook(); book.active.title = "主要"
            person = "背景" if ver == "v1.0.0" else "起源"
            book.create_sheet(person); book.create_sheet("法术大全")
            book["主要"]["A1"] = f"DND 5E card<Test {ver}>"
            book["主要"]["F13"] = 18
            current, maximum = ("M22", "Q22") if ver == "v1.0.0" else ("R22", "V22")
            book["主要"][current] = 7; book["主要"][maximum] = 21
            names = {"主要": "Main", person: "Background" if ver == "v1.0.0" else "Origin", "法术大全": "Spell Compendium"}
            with self.subTest(version=ver):
                actual = self.parsed(WorkbookView(book, names))
                self.assertEqual(actual, self.parsed(book, original=True))
                self.assertEqual(actual["core_stats"]["hp"]["current"], 7)
                self.assertEqual(actual["core_stats"]["hp"]["max"], 21)
                self.assertEqual(actual["meta"]["layout_version"], "v1.0.0" if ver == "v1.0.0" else "v1.0.12")
                with self.assertRaises(ValueError):
                    self.parsed(WorkbookView(book, names, omit=("法术大全",)))
            book.close()

    def test_unknown_legacy_title_is_rejected(self):
        book = Workbook(); book.active.title = "Main"
        book.create_sheet("Background"); book.create_sheet("Spell Compendium")
        book["Main"]["A1"] = "DND card<Test v9.9.9>"
        with self.assertRaisesRegex(ValueError, "Unsupported"):
            self.parsed(book)
        book.close()

    def test_no_shared_layout_mutation_or_cross_import_alias_leak(self):
        before = copy.deepcopy((parser.LAYOUT_V1_0_0, parser.LAYOUT_V1_0_12))
        for version, book in self.books.items():
            parser._select_layout(WorkbookView(book, ENGLISH[version]))
            self.assertEqual(parser._select_layout(book)["sheets"]["main"], "主要")
        self.assertEqual(before, (parser.LAYOUT_V1_0_0, parser.LAYOUT_V1_0_12))


if __name__ == "__main__":
    unittest.main()
