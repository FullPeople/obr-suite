"""Real source-card parser checks using read-only workbook/cell views; no save()."""
from __future__ import annotations

import copy
import hashlib
import importlib.util
import os
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
PARSER = Path(os.environ.get("OBR_SHIELD_PARSER", ROOT.parent / "character-cards-server/parser.py"))
BASELINE = Path(os.environ.get("OBR_SHIELD_BASELINE", ROOT.parent / "_audit/2026-09-08/xlsx-server-shield/parser-before-shield.py"))
PUBLIC = Path(os.environ.get("OBR_SHIELD_PUBLIC", ROOT / "public"))
BEFORE_SHA = "9791888eddaac50e771baaff0cc7e738ef1b5cb18b52a954344ccd2224512c34"
SOURCES = {
    "2014": ("DND5E人物卡_悲灵_弗人_枭熊适配版.xlsx", "94444fda4206d579125418b0007b89c414b86c27d6b98655f81f3f1ce65e8fe6"),
    "2024": ("DND5R人物卡_悲灵_弗人_枭熊适配版.xlsx", "264fc65569e3e80932544af548830ca70f10ee308493c8b0234608fcd01c9b04"),
}


def module(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


class SheetView:
    def __init__(self, source, overrides):
        self.source, self.overrides = source, overrides

    def __getitem__(self, key):
        return SimpleNamespace(value=self.overrides[key]) if key in self.overrides else self.source[key]

    def cell(self, row, column):
        from openpyxl.utils.cell import get_column_letter
        return self[f"{get_column_letter(column)}{row}"]

    def __getattr__(self, key):
        return getattr(self.source, key)


class BookView:
    def __init__(self, source, values=None, main_name="主要"):
        self.sheets = {main_name if name == "主要" else name:
                       SheetView(source[name], values or {} if name == "主要" else {})
                       for name in source.sheetnames}
        self.sheetnames = list(self.sheets)

    def __getitem__(self, name):
        return self.sheets[name]


def stable(data):
    result = copy.deepcopy(data)
    result["meta"].pop("parsed_at")
    return result


class ShieldInputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        assert hashlib.sha256(BASELINE.read_bytes()).hexdigest() == BEFORE_SHA
        cls.old, cls.new = module(BASELINE, "shield_before"), module(PARSER, "shield_after")
        cls.books, cls.paths = {}, {}
        for version, (filename, sha) in SOURCES.items():
            path = PUBLIC / filename
            assert hashlib.sha256(path.read_bytes()).hexdigest() == sha
            cls.paths[version] = path
            cls.books[version] = load_workbook(path, data_only=True)

    @classmethod
    def tearDownClass(cls):
        for version, book in cls.books.items():
            book.close()
            assert hashlib.sha256(cls.paths[version].read_bytes()).hexdigest() == SOURCES[version][1]

    def combat(self, book, old=False):
        mod = self.old if old else self.new
        return mod.parse_combat(book, mod._select_layout(book))

    def parsed(self, book, version, old=False):
        mod = self.old if old else self.new
        with patch.object(mod, "load_workbook", return_value=book):
            return stable(mod.parse_character(self.paths[version]))

    def test_01_reproduced_wrong_coordinate_and_chinese_value(self):
        for version, source in self.books.items():
            for yes in ("是", "Yes"):
                with self.subTest(version=version, yes=yes):
                    view = BookView(source, {"AS40": yes, "AS41": "No"})
                    self.assertFalse(self.combat(view, old=True)["shield"]["equipped"])
                    self.assertTrue(self.combat(view)["shield"]["equipped"])

    def test_02_known_bilingual_headers_and_main_alias(self):
        for version, source in self.books.items():
            for main in ("主要", "Main"):
                for shield in ("盾牌", "Shield"):
                    for ac in ("AC", "Armor Class"):
                        for equipped in ("着装", "Equipped"):
                            with self.subTest(version=version, main=main, shield=shield, ac=ac, equipped=equipped):
                                view = BookView(source, {"AL39": shield, "AQ39": ac, "AS39": equipped,
                                                        "AS40": "Yes", "AS41": "No"}, main)
                                self.assertTrue(self.combat(view)["shield"]["equipped"])
            view = BookView(source, {"AL39": " shield ", "AQ39": " ARMOR CLASS ", "AS39": " equipped ", "AS40": " yes "})
            self.assertTrue(self.combat(view)["shield"]["equipped"])

    def test_03_input_languages_conflicts_and_existing_symbols(self):
        cases = [("是", True), ("否", False), ("Yes", True), ("No", False), (" yes ", True),
                 ("NO", False), (None, False), ("", False), ("有", False), ("无", False),
                 ("O", True), ("Y", True), ("✓", True), (1, True), (True, True),
                 (0, False), (False, False), ("X", False), ("TRUE", True)]
        for version, source in self.books.items():
            for value, expected in cases:
                with self.subTest(version=version, value=value):
                    view = BookView(source, {"AS40": value, "AS41": "No" if expected else "Yes"})
                    self.assertIs(self.combat(view)["shield"]["equipped"], expected)

    def test_04_each_unknown_or_missing_header_keeps_as41_legacy_semantics(self):
        for version, source in self.books.items():
            for ref in ("AL39", "AQ39", "AS39"):
                for header in (None, "", "Custom", "Shield / 盾牌"):
                    for old_value in ("是", "O", "✓", "Y", 1, "YES", "X", None):
                        with self.subTest(version=version, ref=ref, header=header, old_value=old_value):
                            view = BookView(source, {ref: header, "AS40": "Yes", "AS41": old_value})
                            self.assertEqual(self.combat(view), self.combat(view, old=True))

    def test_05_zero_and_empty_bonus_do_not_change_boolean_contract(self):
        for version, source in self.books.items():
            for bonus, parsed in ((0, 0), (None, None), ("", None), ("0", 0), (2, 2)):
                for mark, expected in (("Yes", True), ("No", False)):
                    with self.subTest(version=version, bonus=bonus, mark=mark):
                        view = BookView(source, {"AQ40": bonus, "AS40": mark, "AS41": mark})
                        before, after = self.combat(view, old=True), self.combat(view)
                        self.assertEqual(after, before)
                        self.assertEqual(after["shield"]["ac_bonus"], parsed)
                        self.assertIs(after["shield"]["equipped"], expected)

    def test_06_full_parse_changes_only_equipped_with_cached_totals(self):
        for version, source in self.books.items():
            for mark, expected in (("是", True), ("Yes", True), ("否", False), ("No", False)):
                for main in ("主要", "Main"):
                    with self.subTest(version=version, mark=mark, main=main):
                        view = BookView(source, {"AS40": mark, "AS41": "No" if expected else "Yes"}, main)
                        before, after = self.parsed(view, version, old=True), self.parsed(view, version)
                        self.assertIs(after["combat"]["shield"]["equipped"], expected)
                        before["combat"]["shield"]["equipped"] = expected
                        self.assertEqual(after, before)

    def test_07_overlay_priority_and_cached_ac_weight_stay_unchanged(self):
        for version, source in self.books.items():
            for export, ac in ((source["主要"]["AV1"].value, 10), (None, 12)):
                with self.subTest(version=version, export=export is not None):
                    view = BookView(source, {"AS40": "Yes", "AS41": "No", "D23": 12, "L60": 123, "AV1": export})
                    before, after = self.parsed(view, version, old=True), self.parsed(view, version)
                    self.assertEqual(after["core_stats"]["ac"], ac)
                    self.assertEqual(after["inventory"], before["inventory"])
                    self.assertEqual(after["exports"], before["exports"])
                    before["combat"]["shield"]["equipped"] = True
                    self.assertEqual(after, before)

    def test_08_existing_missing_ac_fallback_consumes_corrected_flag(self):
        for version, source in self.books.items():
            for bonus, delta in ((2, 2), (0, 0), (None, 0), ("", 0)):
                with self.subTest(version=version, bonus=bonus):
                    view = BookView(source, {"AS40": "Yes", "AS41": "No", "AQ40": bonus, "D23": None, "AV1": None})
                    before, after = self.parsed(view, version, old=True), self.parsed(view, version)
                    self.assertEqual(after["core_stats"]["ac"], before["core_stats"]["ac"] + delta)
                    before["combat"]["shield"]["equipped"] = True
                    before["core_stats"]["ac"] += delta
                    self.assertEqual(after, before)

    def test_09_explicit_v100_preserves_legacy_default_and_coordinates(self):
        book = Workbook()
        book.active.title = "主要"
        book.create_sheet("背景"); book.create_sheet("法术大全")
        book["主要"]["A1"] = "DND 5E 人物卡<悲灵v1.0.0>"
        for ref, value in {"AL39": "Shield", "AQ39": "AC", "AS39": "Equipped", "AS40": "No", "AS41": "No", "AO31": 2}.items():
            book["主要"][ref] = value
        self.assertEqual(self.combat(book), self.combat(book, old=True))
        self.assertTrue(self.combat(book)["shield"]["equipped"])
        self.assertEqual(self.combat(book)["shield"]["ac_bonus"], 2)
        book.close()
        # Even a mixed modern signature must not silently override an explicit old title.
        view = BookView(self.books["2014"], {"A1": "DND 5E 人物卡<悲灵v1.0.0>", "AS40": "Yes", "AS41": "No"})
        self.assertEqual(self.combat(view), self.combat(view, old=True))

    def test_10_missing_version_evidence_and_ambiguous_main_remain_rejected(self):
        view = BookView(self.books["2014"], {**{f"{c}2": None for c in "EFGHIJKLMN"}, "AV1": None,
                                                   "A1": "Unknown card", "AS40": "Yes"})
        with self.assertRaises(ValueError): self.new._select_layout(view)
        view = BookView(self.books["2014"])
        view.sheets["Main"] = view.sheets["主要"]; view.sheetnames.append("Main")
        with self.assertRaises(ValueError): self.new._select_layout(view)

    def test_11_layout_constants_and_subsequent_imports_are_not_mutated(self):
        before = copy.deepcopy(self.new.LAYOUT_V1_0_12)
        first = BookView(self.books["2014"], {"AS40": "Yes", "AS41": "No"})
        self.assertTrue(self.combat(first)["shield"]["equipped"])
        second = BookView(self.books["2014"], {"AL39": "Unknown", "AS40": "Yes", "AS41": "No"})
        self.assertEqual(self.combat(second), self.combat(second, old=True))
        self.assertEqual(self.new.LAYOUT_V1_0_12, before)
        self.assertEqual(before["main"]["shield_equipped"], "AS41")

    def test_12_original_full_parse_is_unchanged(self):
        for version, source in self.books.items():
            self.assertEqual(self.parsed(source, version), self.parsed(source, version, old=True))


if __name__ == "__main__":
    unittest.main(verbosity=2)
