"""Exercise actual parsers on immutable originals and in-memory input views.

Legacy cases exercise the actual old coordinate configuration with seeded values;
they do not claim access to an original v1.0.0 workbook or a spreadsheet engine.
"""
import ast
import copy
import hashlib
import importlib.util
import os
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

if not __debug__:
    raise RuntimeError('Verification requires assertions; do not use -O')

PARSER=Path(os.environ['OBR_ABILITY_INPUTS_PARSER'])
BASELINE=Path(os.environ['OBR_ABILITY_INPUTS_BASELINE'])
PUBLIC=Path(os.environ['OBR_ABILITY_INPUTS_PUBLIC'])
SOURCES={
    '2014':('DND5E人物卡_悲灵_弗人_枭熊适配版.xlsx','94444fda4206d579125418b0007b89c414b86c27d6b98655f81f3f1ce65e8fe6'),
    '2024':('DND5R人物卡_悲灵_弗人_枭熊适配版.xlsx','264fc65569e3e80932544af548830ca70f10ee308493c8b0234608fcd01c9b04'),
}
PAIRS=(('智力','Intelligence'),('感知','Wisdom'),('魅力','Charisma'))
sha=lambda b:hashlib.sha256(b).hexdigest()

def module(path,name):
    spec=importlib.util.spec_from_file_location(name,path)
    result=importlib.util.module_from_spec(spec);spec.loader.exec_module(result);return result

class SheetView:
    def __init__(self,source,values):self.source,self.values=source,values
    def __getitem__(self,ref):return SimpleNamespace(value=self.values[ref]) if ref in self.values else self.source[ref]
    def cell(self,row,column):return self[f'{get_column_letter(column)}{row}']
    def __getattr__(self,key):return getattr(self.source,key)

class BookView:
    def __init__(self,source,values,english=False):
        self.sheetnames=['Main' if english and n=='主要' else n for n in source.sheetnames]
        self.sheets=dict(zip(self.sheetnames,[SheetView(source[n],values) if n=='主要' else source[n] for n in source.sheetnames]))
    def __getitem__(self,key):return self.sheets[key]

class AbilityInputsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.before=BASELINE.read_bytes();cls.after=PARSER.read_bytes()
        assert sha(cls.before)=='9272d4edda9cab30aa1acf9ed7b06b9d9192ff248144654307f33e62c553b7d4'
        cls.old=module(BASELINE,'ability_before');cls.new=module(PARSER,'ability_after')
        cls.books={}
        for version,(name,digest) in SOURCES.items():
            assert sha((PUBLIC/name).read_bytes())==digest
            cls.books[version]=load_workbook(PUBLIC/name,data_only=True)

    @classmethod
    def tearDownClass(cls):
        for version,book in cls.books.items():
            book.close();name,digest=SOURCES[version];assert sha((PUBLIC/name).read_bytes())==digest
        assert BASELINE.read_bytes()==cls.before and PARSER.read_bytes()==cls.after

    def layouts(self):
        for version,book in self.books.items():
            for english in (False,True):
                view=BookView(book,{},english)
                yield version,english,self.new._select_layout(view)
        legacy=copy.deepcopy(self.new.LAYOUT_V1_0_0)
        legacy['sheets']={**self.new._select_layout(self.books['2014'])['sheets']}
        yield '2014',False,legacy

    def values(self,layout,ability):
        m=layout['main']
        return {m['spell_active_ability_cell']:ability,
                **{f"{m['spell_dc_col']}{row}":100+row for row in m['spell_ability_to_row'].values()},
                **{f"{m['spell_atk_col']}{row}":'Attack '+str(row) for row in m['spell_ability_to_row'].values()}}

    def parsed(self,mod,version,english,layout,ability):
        return mod.parse_spellcasting(BookView(self.books[version],self.values(layout,ability),english),layout)

    def test_01_only_spellcasting_function_changes(self):
        before,after=ast.parse(self.before),ast.parse(self.after)
        self.assertEqual(len(before.body),len(after.body))
        changed=[getattr(a,'name',None) for a,b in zip(before.body,after.body) if ast.dump(a)!=ast.dump(b)]
        self.assertEqual(changed,['parse_spellcasting'])
        old='    active_row = abil_to_row.get(active_abil or "", default_row)'
        new='''    ability_key = active_abil or ""
    if ability_key not in abil_to_row:
        ability_key = {"intelligence": "智力", "wisdom": "感知", "charisma": "魅力"}.get(
            ability_key.casefold(), ability_key)
    active_row = abil_to_row.get(ability_key, default_row)'''
        self.assertEqual(self.before.decode().replace(old,new).encode(),self.after)

    def test_02_actual_original_full_parse_unchanged(self):
        for version,(name,_) in SOURCES.items():
            results=[]
            for mod in (self.old,self.new):
                with patch.object(mod,'load_workbook',return_value=self.books[version]):data=mod.parse_character(PUBLIC/name)
                data['meta'].pop('parsed_at');results.append(data)
            self.assertEqual(*results)

    def test_03_chinese_empty_unknown_and_physical_abilities_preserved(self):
        for version,english,layout in self.layouts():
            for value in ['智力','感知','魅力','力量','敏捷','体质','Strength','Dexterity','Constitution','Unknown ability','',None,123]:
                with self.subTest(version=version,english=english,input=value,cell=layout['main']['spell_active_ability_cell']):
                    self.assertEqual(self.parsed(self.old,version,english,layout,value),self.parsed(self.new,version,english,layout,value))

    def test_04_english_aliases_select_correct_actual_dc_attack_rows(self):
        for version,english,layout in self.layouts():
            for zh,en in PAIRS:
                for value in (en,en.lower(),en.upper(),' '+en+' '):
                    with self.subTest(version=version,english=english,input=value,cell=layout['main']['spell_active_ability_cell']):
                        baseline=self.parsed(self.old,version,english,layout,zh)
                        actual=self.parsed(self.new,version,english,layout,value)
                        self.assertEqual(actual['spellcasting_ability'],value.strip())
                        row=layout['main']['spell_ability_to_row'][zh]
                        self.assertEqual(actual['save_dc'],100+row)
                        self.assertEqual(actual['attack_bonus'],'Attack '+str(row))
                        actual['spellcasting_ability']=zh;self.assertEqual(actual,baseline)

    def test_05_existing_explicit_english_mapping_takes_precedence(self):
        layout=copy.deepcopy(self.new._select_layout(self.books['2014']))
        layout['main']['spell_ability_to_row']['Intelligence']=70
        for mod in (self.old,self.new):
            data=self.parsed(mod,'2014',False,layout,'Intelligence');self.assertEqual(data['save_dc'],170)

    def test_06_no_layout_or_module_configuration_mutation(self):
        originals={key:copy.deepcopy(value) for key,value in vars(self.new).items() if key.startswith('LAYOUT_')}
        for version,english,layout in self.layouts():
            original=copy.deepcopy(layout)
            for _,en in PAIRS:self.parsed(self.new,version,english,layout,en)
            self.assertEqual(original,layout)
        self.assertEqual(originals,{key:value for key,value in vars(self.new).items() if key.startswith('LAYOUT_')})

    def test_07_whole_character_keeps_display_and_unrelated_fields(self):
        for version,book in self.books.items():
            layout=self.new._select_layout(book)
            for zh,en in PAIRS:
                outputs=[]
                for mod,ability in [(self.old,zh),(self.new,en)]:
                    view=BookView(book,self.values(layout,ability),True)
                    with patch.object(mod,'load_workbook',return_value=view):result=mod.parse_character(PUBLIC/SOURCES[version][0])
                    result['meta'].pop('parsed_at');outputs.append(result)
                self.assertEqual(outputs[1]['spellcasting']['spellcasting_ability'],en)
                self.assertEqual(outputs[1]['core_stats']['dc_ability'],en)
                outputs[1]['spellcasting']['spellcasting_ability']=zh
                outputs[1]['core_stats']['dc_ability']=zh
                self.assertEqual(*outputs)

if __name__=='__main__':unittest.main(verbosity=2)
