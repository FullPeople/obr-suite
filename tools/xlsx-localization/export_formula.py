"""Plan, but never write XLSX, a safe replacement for the two original AV1s.

This is a deliberately restricted formula compiler, not an Excel interpreter or
schema migration. It accepts the originals' CONCAT/IF/ISNUMBER/TEXTJOIN grammar,
proves JSON string boundaries through each branch, and escapes only data inside
strings. Numeric guards, blank-row filters, order, types and source references
(including AS41) survive unchanged. Unsupported source syntax fails closed.

All new cells live on one explicitly declared, new hidden Export sheet. AV1
keeps the original machine-header literals. Length cells are calculated before
value cells, so overflow yields an explicit non-JSON error, never a prefix or a
previous cache. Generated cells have no cached values. evaluate_plan is a small
reference model over supplied inputs; it does NOT recalculate their formulas or
prove Excel/WPS, locale coercion, internal token bytes, or Web Excel compatibility.

Desktop legacy functions only: IF, ISNUMBER, TEXTJOIN (original _xlfn spelling),
SUBSTITUTE, CHAR, LEN, SUM. CHAR(0) is never generated. U+0001..001F are escaped;
NUL, lone surrogates and XML-forbidden noncharacters are rejected as inputs.
CHAR for all these controls still needs actual target Excel/WPS acceptance;
Excel for the web documents a narrower CHAR subset. No CLEAN or normalization.

References: https://support.microsoft.com/en-us/excel/excel-specifications-and-limits
https://support.microsoft.com/en-us/excel/functions/substitute-function
https://support.microsoft.com/en-us/excel/functions/char-function

Run with the bundled Python and -B; see export_formula_selftest.py. CLI:
  export_formula.py --source original.xlsx --output audit/plan.json
  Optional --sheet-map '{"主要":"Main","背景":"Background"}' describes
  a separately approved rename prerequisite; this tool does not perform it.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
import hashlib
import json
import math
from pathlib import Path
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET

from openpyxl.formula.tokenizer import Tokenizer

MAX_FORMULA = 8192
MAX_TEXT = 32767
MAX_ROW = 1048576
MAX_COLUMN = 16384
OVERFLOW = "ERROR: OBR_EXPORT_OUTPUT_TOO_LONG (maximum 32767 characters); no JSON exported."
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


class Rejected(ValueError):
    """Unsupported or invalid input, never a plausible fallback export."""


def require(ok, message):
    if not ok:
        raise Rejected(message)


def utf16len(s):
    return len(s.encode("utf-16-le")) // 2


def valid_text(value):
    require(value is None or type(value) in (str, int, float, bool) or isinstance(value, ExcelError),
            "Unsupported scalar cell value")
    if isinstance(value, str):
        require(not any(ord(c) == 0 or 0xD800 <= ord(c) <= 0xDFFF or ord(c) in (0xFFFE, 0xFFFF)
                        for c in value), "Unrepresentable text: NUL, lone surrogate or XML noncharacter")
        require(utf16len(value) <= MAX_TEXT, "Source cell exceeds 32767 UTF-16 code units")
    elif isinstance(value, float):
        require(math.isfinite(value), "Non-finite numeric input")


def as_text(value):
    if isinstance(value, ExcelError):
        raise Rejected(f"Source formula error: {value.value}")
    if value is None:
        return ""
    if value is True:
        return "TRUE"
    if value is False:
        return "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


@dataclass(frozen=True)
class ExcelError:
    value: str


@dataclass(frozen=True)
class Node:
    op: str
    value: object = None
    args: tuple = ()


def lit(value):
    return Node("lit", value)


class Formula:
    """Parse only the scalar grammar needed by originals and generated plans."""
    def __init__(self, formula):
        require(isinstance(formula, str) and bool(formula), "Missing formula")
        self.tokens = [t for t in Tokenizer(formula if formula.startswith("=") else "=" + formula).items
                       if t.type != "WHITE-SPACE"]
        self.pos = 0
        self.tree = self.expr()
        require(self.pos == len(self.tokens), "Trailing/unsupported formula tokens")

    def expr(self, minimum=0):
        left = self.atom()
        precedence = {"=": 1, ">": 1, "&": 2, "+": 3, "-": 3, "*": 4}
        while self.pos < len(self.tokens):
            token = self.tokens[self.pos]
            level = precedence.get(token.value, -1)
            if token.type != "OPERATOR-INFIX" or level < minimum:
                break
            self.pos += 1
            right = self.expr(level + 1)
            # Flat concatenation avoids recursion on the originals' long spine.
            if token.value == "&" and left.op == "&":
                left = Node("&", args=left.args + (right,))
            else:
                left = Node(token.value, args=(left, right))
        return left

    def atom(self):
        require(self.pos < len(self.tokens), "Incomplete formula")
        token = self.tokens[self.pos]
        self.pos += 1
        if token.type == "FUNC" and token.subtype == "OPEN":
            args = []
            while True:
                args.append(self.expr())
                require(self.pos < len(self.tokens), "Unclosed function")
                end = self.tokens[self.pos]
                self.pos += 1
                if end.type == "FUNC" and end.subtype == "CLOSE":
                    break
                require(end.type == "SEP" and end.subtype == "ARG", "Unsupported function separator")
            return Node("fn", token.value[:-1], tuple(args))
        if token.type == "PAREN" and token.subtype == "OPEN":
            result = self.expr()
            require(self.pos < len(self.tokens) and self.tokens[self.pos].type == "PAREN"
                    and self.tokens[self.pos].subtype == "CLOSE", "Unclosed parentheses")
            self.pos += 1
            return result
        require(token.type == "OPERAND", f"Unsupported formula operand: {token.value}")
        if token.subtype == "TEXT":
            return lit(token.value[1:-1].replace('""', '"'))
        if token.subtype == "NUMBER":
            return lit(float(token.value) if any(c in token.value.lower() for c in ".e") else int(token.value))
        if token.subtype == "LOGICAL":
            return lit(token.value.upper() == "TRUE")
        if token.subtype == "RANGE":
            return Node("ref", token.value)
        raise Rejected(f"Unsupported operand subtype: {token.subtype}")


def reference(value, default_sheet):
    if "!" in value:
        sheet, cell = value.rsplit("!", 1)
        if sheet.startswith("'") and sheet.endswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
    else:
        sheet, cell = default_sheet, value
    require(bool(sheet) and not any(c in sheet for c in "[]"), "External or empty sheet reference")
    require(re.fullmatch(r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}", cell), f"Not a scalar A1 reference: {value}")
    normalized = cell.replace("$", "").upper()
    column, row = re.fullmatch(r"([A-Z]+)([0-9]+)", normalized).groups()
    column_number = 0
    for letter in column:
        column_number = column_number * 26 + ord(letter) - ord("A") + 1
    require(column_number <= MAX_COLUMN and int(row) <= MAX_ROW,
            f"A1 reference outside Excel worksheet bounds: {value}")
    return sheet, normalized


def qualified(sheet, cell):
    col, row = re.fullmatch(r"([A-Z]+)([0-9]+)", cell).groups()
    return "'" + sheet.replace("'", "''") + "'!$" + col + "$" + row


def quoted(value):
    valid_text(value)
    return '"' + value.replace('"', '""') + '"'


def render(node, main_sheet, sheet_map):
    if node.op == "lit":
        return quoted(node.value) if isinstance(node.value, str) else as_text(node.value)
    if node.op == "ref":
        sheet, cell = reference(node.value, main_sheet)
        return qualified(sheet_map.get(sheet, sheet), cell)
    if node.op == "fn":
        return str(node.value) + "(" + ",".join(render(a, main_sheet, sheet_map) for a in node.args) + ")"
    return "(" + node.op.join(render(a, main_sheet, sheet_map) for a in node.args) + ")"


def walk(node):
    yield node
    for child in node.args:
        yield from walk(child)


def evaluate(tree, values, main_sheet, resolve=None):
    def ev(n):
        if n.op == "lit":
            return n.value
        if n.op == "ref":
            key = reference(n.value, main_sheet)
            return resolve(key) if resolve else values.get(key)
        if n.op == "fn":
            name = str(n.value).removeprefix("_xlfn.").upper()
            a = n.args
            if name == "IF":
                return ev(a[1] if ev(a[0]) else a[2])
            if name == "ISNUMBER":
                v = ev(a[0])
                return type(v) in (int, float)
            if name == "TEXTJOIN":
                delimiter, skip = as_text(ev(a[0])), ev(a[1])
                parts = [as_text(ev(x)) for x in a[2:]]
                return delimiter.join(x for x in parts if x or not skip)
            if name == "SUBSTITUTE":
                return as_text(ev(a[0])).replace(as_text(ev(a[1])), as_text(ev(a[2])))
            if name == "CHAR":
                code = ev(a[0])
                require(type(code) is int and 1 <= code <= 127, "Model only supports portable ASCII CHAR(1..127)")
                return chr(code)
            if name == "LEN":
                return utf16len(as_text(ev(a[0])))
            if name == "SUM":
                return sum(ev(x) for x in a)
            raise Rejected(f"Unsupported model function: {name}")
        vals = [ev(a) for a in n.args]
        if n.op == "&":
            return "".join(as_text(v) for v in vals)
        if n.op == "=":
            x, y = vals
            if isinstance(x, ExcelError) or isinstance(y, ExcelError):
                raise Rejected("Error in formula condition")
            x, y = ("" if x is None else x), ("" if y is None else y)
            if isinstance(x, str) and isinstance(y, str):
                return x.casefold() == y.casefold()
            return type(x) is type(y) and x == y or type(x) in (int, float) and type(y) in (int, float) and x == y
        if n.op == ">": return vals[0] > vals[1]
        if n.op == "+": return vals[0] + vals[1]
        if n.op == "-": return vals[0] - vals[1]
        if n.op == "*": return vals[0] * vals[1]
        raise Rejected(f"Unsupported model operator: {n.op}")
    return ev(tree)


def boundary(text, in_string):
    """Static JSON quote context; reject escape sequences split across nodes."""
    escaped = False
    for char in text:
        if escaped:
            require(in_string, "JSON escape outside string")
            escaped = False
        elif char == "\\" and in_string:
            escaped = True
        elif char == '"':
            in_string = not in_string
    require(not escaped, "JSON escape crosses a dynamic boundary")
    return in_string


@dataclass(frozen=True)
class Part:
    value: str  # formula expression
    length: str  # formula expression, measured before concatenation


class Compiler:
    def __init__(self, main, names, sheet_map=None, export_sheet="Export"):
        self.main, self.names = main, names
        self.mapping = dict(sheet_map or {})
        require(set(self.mapping) <= set(names), "Rename prerequisite names a missing sheet")
        targets = [self.mapping.get(name, name) for name in names]
        for name in [*targets, export_sheet]:
            require(isinstance(name, str), "Sheet names must be strings")
            valid_text(name)
            # Data cells keep CR/LF/tab for JSON escaping. Names are instead
            # used directly in worksheet metadata and qualified references.
            require(not any(ord(c) < 32 for c in name), "Control character in sheet name")
        require(len({s.casefold() for s in targets}) == len(targets), "Sheet rename collision")
        for name in [*targets, export_sheet]:
            require(bool(name) and utf16len(name) <= 31 and not any(c in name for c in ":\\/?*[]")
                    and not name.startswith("'") and not name.endswith("'"), "Invalid sheet name")
        require(export_sheet.casefold() not in {x.casefold() for x in targets}, "Export sheet already exists")
        require(self.mapping.get(main, main) in ("主要", "Main"), "Main sheet must remain identifiable")
        self.export, self.cells, self.escaped = export_sheet, [], {}

    def cell(self, expr, purpose):
        formula = "=" + expr
        require(utf16len(formula) <= MAX_FORMULA, f"Formula too long: {purpose}")
        # Conservative textual bytes only, NOT Excel's internal token size.
        require(len(formula.encode("utf-16-le")) <= 16384, "Formula text exceeds byte budget")
        cell = "A" + str(len(self.cells) + 1)
        self.cells.append({"cell": cell, "formula": formula, "purpose": purpose})
        return qualified(self.export, cell)

    def part(self, expr, length, purpose):
        length_ref = self.cell(length, purpose + ": length")
        value_ref = self.cell(f"IF({length_ref}>{MAX_TEXT},{quoted(OVERFLOW)},{expr})", purpose + ": value")
        return Part(value_ref, length_ref)

    def escape(self, node):
        key = reference(node.value, self.main)
        require(key[0] in self.names, "Missing referenced sheet")
        if key in self.escaped:
            return self.escaped[key]
        source = render(node, self.main, self.mapping) + '&""'
        # Backslash first; all subsequently introduced escapes are protected.
        replacements = [(92, "\\\\"), (34, '\\"')]
        replacements += [(n, {8: "\\b", 9: "\\t", 10: "\\n", 12: "\\f", 13: "\\r"}.get(n, f"\\u{n:04x}"))
                         for n in range(1, 32)]
        raw = self.cell(source, f"text input {key[0]}!{key[1]}")
        changes, escaped = [], raw
        for code, replacement in replacements:
            char = f"CHAR({code})"
            count = f"(LEN({raw})-LEN(SUBSTITUTE({raw},{char},\"\")))"
            changes.append(f"{count}*{len(replacement)-1}")
            escaped = f"SUBSTITUTE({escaped},{char},{quoted(replacement)})"
        result = self.part(escaped, f"SUM(LEN({raw})," + ",".join(changes) + ")", f"JSON escape {key[0]}!{key[1]}")
        self.escaped[key] = result
        return result

    def scalar(self, node):
        """Admit only original numeric guard or a blank comparison condition."""
        if node.op == "ref":
            require(reference(node.value, self.main)[0] in self.names, "Missing referenced sheet")
        elif node.op == "lit":
            pass
        elif node.op == "=" and len(node.args) == 2:
            for a in node.args: self.scalar(a)
        elif node.op == "fn" and node.value in ("IF", "ISNUMBER"):
            require(len(node.args) == (3 if node.value == "IF" else 1), "Wrong scalar arity")
            for a in node.args: self.scalar(a)
        else:
            raise Rejected("Unsupported scalar control expression")
        return render(node, self.main, self.mapping)

    def compile(self, node, state=False):
        if node.op == "lit":
            if isinstance(node.value, str):
                return Part(quoted(node.value), str(utf16len(node.value))), boundary(node.value, state)
            require(not state and type(node.value) in (int, float), "Unsupported literal in JSON output")
            expr = as_text(node.value)
            return Part(expr, str(len(expr))), state
        if node.op == "ref":
            require(state, "Unguarded dynamic output outside JSON string")
            return self.escape(node), state
        if node.op == "&":
            parts = []
            for a in node.args:
                part, state = self.compile(a, state)
                parts.append(part)
            return self.join(parts), state
        if node.op == "fn" and node.value == "IF":
            require(len(node.args) == 3, "IF arity")
            cond = self.scalar(node.args[0])
            # Retain precisely the original typed number-or-zero path.
            a = node.args
            numeric = (not state and a[0].op == "fn" and a[0].value == "ISNUMBER"
                       and len(a[0].args) == 1 and a[0].args[0].op == "ref"
                       and a[1] == a[0].args[0] and a[2] == lit(0))
            if numeric:
                expr = self.scalar(node)
                return self.part(expr, f"LEN({expr}&\"\")", "original numeric guard"), state
            yes, yes_state = self.compile(a[1], state)
            no, no_state = self.compile(a[2], state)
            require(yes_state == no_state, "IF branches disagree on JSON string boundary")
            return self.part(f"IF({cond},{yes.value},{no.value})",
                             f"IF({cond},{yes.length},{no.length})", "original IF branch"), yes_state
        if node.op == "fn" and node.value in ("TEXTJOIN", "_xlfn.TEXTJOIN"):
            require(not state and len(node.args) >= 3, "TEXTJOIN must join complete JSON fragments")
            delimiter, skip, *items = node.args
            require(delimiter == lit(",") and skip == lit(True), "Unsupported TEXTJOIN delimiter/blank rule")
            parts = []
            for a in items:
                p, end = self.compile(a, False)
                require(not end, "TEXTJOIN item has open JSON string")
                parts.append(p)
            count = self.cell("SUM(" + ",".join(f"IF({p.length}=0,0,1)" for p in parts) + ")", "nonempty TEXTJOIN items")
            length = "SUM(" + ",".join(p.length for p in parts) + f")+IF({count}=0,0,{count}-1)"
            expr = str(node.value) + '(\",\",TRUE,' + ",".join(p.value for p in parts) + ")"
            return self.part(expr, length, "original TEXTJOIN"), False
        raise Rejected(f"Unsupported JSON-output AST: {node.op} {node.value}")

    def join(self, parts):
        require(bool(parts), "Empty concatenation")
        # Bounded groups preserve ordering and keep every expression short.
        if len(parts) > 20:
            return self.join([self.join(parts[i:i+20]) for i in range(0, len(parts), 20)])
        return self.part("&".join(p.value for p in parts),
                         "SUM(" + ",".join(p.length for p in parts) + ")", "ordered JSON concatenation")


def make_plan(formula, *, main_sheet, sheet_names, sheet_map=None, export_sheet="Export"):
    parsed = Formula(formula).tree
    require(parsed.op == "&" and len(parsed.args) >= 4, "Unrecognized AV1 header structure")
    head = parsed.args[:3]
    require(head[0].op == head[2].op == "lit" and head[1] == Node("ref", "A1"), "Unrecognized AV1 template header")
    require('"schema":"obr-suite-card/v1"' in head[0].value, "Missing original schema literal")
    rules = re.findall(r'"ruleset":"(5E2014|5E2024)"', head[2].value)
    require(len(rules) == 1, "Missing or ambiguous original ruleset literal")
    compiler = Compiler(main_sheet, sheet_names, sheet_map, export_sheet)
    parts, state = [], False
    for node in parsed.args:
        part, state = compiler.compile(node, state)
        parts.append(part)
    require(not state, "AV1 ends inside a JSON string")
    tail = compiler.join(parts[3:])
    all_parts = parts[:3] + [tail]
    length = compiler.cell("SUM(" + ",".join(p.length for p in all_parts) + ")", "complete AV1 output length")
    final = "=" + f"IF({length}>{MAX_TEXT},{quoted(OVERFLOW)}," + "&".join(p.value for p in all_parts) + ")"
    require(utf16len(final) <= MAX_FORMULA, "AV1 formula exceeds length limit")
    require('""schema"":""obr-suite-card/v1""' in final and '""ruleset"":""' + rules[0] + '""' in final,
            "AV1 lost the frontend machine-identification literals")
    mapping = dict(sheet_map or {})
    return {"schema": "obr-xlsx-export-formula-plan/v1", "ruleset": rules[0], "release_ready": False,
            "source_formula_sha256": hashlib.sha256(formula.lstrip("=").encode()).hexdigest(),
            "rename_prerequisites": mapping,
            "add_sheet": {"name": export_sheet, "state": "hidden", "position": "append", "cells": compiler.cells},
            "replace_cell": {"sheet": mapping.get(main_sheet, main_sheet), "cell": "AV1", "formula": final},
            "output_length_cell": length,
            "escaped_sources": [{"sheet": s, "cell": c} for s, c in compiler.escaped],
            "limits": {"formula_utf16_units_including_equals": MAX_FORMULA, "output_utf16_units": MAX_TEXT,
                       "overflow_result": OVERFLOW, "input_rejected": ["NUL", "lone surrogate", "U+FFFE", "U+FFFF"]},
            "preserved_contract": {"schema": "obr-suite-card/v1", "shield_equipped_source": "AS41 unchanged",
                "web_import": "Existing Web Import/网页导入 references to main AV1 stay valid; rename prerequisites are separate.",
                "no_cache_values": True, "no_schema_0_3_migration": True},
            "limitations": ["Reference-model evaluation is not Excel/WPS recalculation or fresh precedent caches.",
                "CHAR(1..31) needs desktop Excel/WPS acceptance; Web Excel supports a narrower CHAR subset.",
                "UTF-16 length checks are conservative for newer Unicode-aware Excel; internal formula token bytes need Excel validation.",
                "No XLSX writer, no cached-value generation, and no importer or dependency migration are included."]}


def evaluate_plan(plan, values):
    for value in values.values(): valid_text(value)
    sheet = plan["add_sheet"]["name"]
    formulas = {(sheet, c["cell"]): Formula(c["formula"]).tree for c in plan["add_sheet"]["cells"]}
    dest = plan["replace_cell"]
    formulas[dest["sheet"], "AV1"] = Formula(dest["formula"]).tree
    memo, visiting = {}, set()
    def resolve(key):
        if key in memo: return memo[key]
        if key not in formulas: return values.get(key)
        require(key not in visiting, "Cycle in export plan")
        visiting.add(key)
        value = evaluate(formulas[key], values, key[0], resolve)
        if isinstance(value, str): require(utf16len(value) <= MAX_TEXT, "Intermediate text exceeds Excel cell capacity")
        visiting.remove(key)
        memo[key] = value
        return value
    result = resolve((dest["sheet"], "AV1"))
    if result == OVERFLOW:
        raise Rejected(OVERFLOW)
    json.loads(result)  # Invalid output is never reported as a successful export.
    return result


def read_original(path):
    """Read the actual package and existing caches; no workbook save path exists."""
    path = Path(path)
    data = path.read_bytes()
    with zipfile.ZipFile(path) as z:
        require(len(z.namelist()) == len(set(z.namelist())), "Duplicate ZIP part")
        q = lambda name: "{" + NS + "}" + name
        book = ET.fromstring(z.read("xl/workbook.xml"))
        rels = {r.get("Id"): r for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
        sheets = {}
        for s in book.find(q("sheets")):
            r = rels[s.get("{" + RNS + "}id")]
            require(r.get("TargetMode") != "External", "External worksheet target")
            target = r.get("Target")
            part = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
            require(part.startswith("xl/worksheets/") and part in z.namelist(), "Unsafe worksheet path")
            sheets[s.get("name")] = part
        mains = [s for s in sheets if s in ("主要", "Main")]
        require(len(mains) == 1, "Require one identifiable main sheet")
        main = mains[0]
        roots = {main: ET.fromstring(z.read(sheets[main]))}
        av1 = roots[main].find(f".//{q('c')}[@r='AV1']")
        require(av1 is not None and av1.find(q("f")) is not None, "Main AV1 has no formula")
        formula = av1.find(q("f")).text
        refs = {reference(n.value, main) for n in walk(Formula(formula).tree) if n.op == "ref"}
        def unescape(s):
            # Single pass preserves the literal _x000A_ encoded as _x005F_x000A_.
            decoded = re.sub(r"_x([0-9A-Fa-f]{4})_", lambda m: chr(int(m[1], 16)), s)
            return decoded.encode("utf-16-le", "surrogatepass").decode("utf-16-le")
        def rich(e): return unescape("".join(t.text or "" for t in e.iter(q("t"))))
        shared = [rich(si) for si in ET.fromstring(z.read("xl/sharedStrings.xml"))] if "xl/sharedStrings.xml" in z.namelist() else []
        def value(cell):
            if cell is None: return None
            if cell.get("t") == "inlineStr": return rich(cell)
            v = cell.find(q("v"))
            if v is None or v.text is None: return None
            if cell.get("t") == "s":
                index = v.text.strip()
                require(len(index) <= 10 and re.fullmatch(r"[0-9]+", index) is not None,
                        "Invalid shared-string index")
                require(int(index) < len(shared), "Shared-string index outside table")
                return shared[int(index)]
            if cell.get("t") == "b": return v.text == "1"
            if cell.get("t") == "e": return ExcelError(v.text)
            if cell.get("t") == "str": return unescape(v.text)
            return float(v.text) if any(c in v.text.lower() for c in ".e") else int(v.text)
        values = {}
        for sheet in {s for s, _ in refs}:
            require(sheet in sheets, "Missing source sheet")
            root = roots.setdefault(sheet, ET.fromstring(z.read(sheets[sheet])))
            cells = {c.get("r"): c for c in root.iter(q("c"))}
            for s, cell in refs:
                if s == sheet: values[s, cell] = value(cells.get(cell))
        cached = value(av1)
        require(isinstance(cached, str), "Original has no existing AV1 text cache for comparison")
        return {"path": str(path.resolve()), "sha256": hashlib.sha256(data).hexdigest(), "formula": formula,
                "main_sheet": main, "sheet_names": list(sheets), "values": values,
                "cached_json": json.loads(cached), "av1_part": sheets[main]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sheet-map", default="{}")
    args = parser.parse_args()
    require(args.output.suffix.lower() == ".json" and args.output.resolve() != args.source.resolve(), "Only a separate JSON plan may be written")
    require(not args.output.exists(), "Refusing to overwrite an existing plan")
    snapshot = read_original(args.source)
    plan = make_plan(snapshot["formula"], main_sheet=snapshot["main_sheet"], sheet_names=snapshot["sheet_names"], sheet_map=json.loads(args.sheet_map))
    baseline = json.loads(evaluate(Formula(snapshot["formula"]).tree, snapshot["values"], snapshot["main_sheet"]))
    require(baseline == snapshot["cached_json"], "Original reference model disagrees with existing AV1 JSON cache")
    renamed_values = {(plan["rename_prerequisites"].get(s, s), c): v for (s, c), v in snapshot["values"].items()}
    require(json.loads(evaluate_plan(plan, renamed_values)) == baseline, "Generated baseline differs from original AV1")
    plan["source_workbook_sha256"] = snapshot["sha256"]
    plan["source_av1_part"] = snapshot["av1_part"]
    plan["reference_model_baseline_equals_existing_json"] = True
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(plan, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"plan": str(args.output), "new_formula_cells": len(plan["add_sheet"]["cells"]), "xlsx_written": False}))


if __name__ == "__main__":
    main()
